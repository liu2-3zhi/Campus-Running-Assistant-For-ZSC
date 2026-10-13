"""多账号导入格式兼容与学号/密码校验回归测试。

覆盖:
    1. “副本账号导入模板”（标题行 + 序号/账号/密码/标记）必须按表头映射，
       序号列只用于人工对行，导入时忽略；
    2. 兼容 只有表头 / 完全没有表头 / 无表头但带序号列 的 xlsx、xls、csv；
    3. 学号必须是 ≥4 位纯数字、密码不得包含中文；
    4. 名单式误导入（序号/姓名）必须被整体拦下并完整上报，不能再生成
       1.ini..21.ini、序号.ini 这类账号；
    5. 下载模板与“副本”同款，且能直接回灌导入；
    6. 加载/保存层过滤历史坏数据，避免坏账号被再次写回。
"""

import base64
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl
import xlrd

import api_runtime
import main as main_module
from api_runtime import _normalize_import_cell

try:
    import xlwt

    _HAS_XLWT = True
except Exception:  # pragma: no cover - 环境缺 xlwt 时跳过 .xls 生成测试
    xlwt = None
    _HAS_XLWT = False

VALID_A = "2024050101059"
VALID_B = "2024050104038"
FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"
LEGACY_TEMPLATE_FILENAME = "legacy_accounts_template.xlsx"
CANONICAL_TEMPLATE_FILENAME = "school_accounts_template.xlsx"


class _FakeAccount:
    def __init__(self, username):
        self.username = username
        self.password = "pw"
        self.tag = ""
        self.params = {}
        self.device_ua = "ua"
        self.user_data = mock.Mock()
        self.is_first_login_verified = True


class _ImportTestBase(unittest.TestCase):
    def setUp(self):
        # 生产环境的 openpyxl 由 main.py 启动时动态导入；测试里直接注入。
        self._openpyxl_patch = mock.patch.object(
            main_module, "openpyxl", openpyxl, create=True
        )
        self._openpyxl_patch.start()
        self.addCleanup(self._openpyxl_patch.stop)
        self._xlrd_patch = mock.patch.object(
            main_module, "xlrd", xlrd, create=True
        )
        self._xlrd_patch.start()
        self.addCleanup(self._xlrd_patch.stop)

    def _make_api(self, existing=()):
        api = object.__new__(main_module.Api)
        api.accounts = {u: _FakeAccount(u) for u in existing}
        api.auth_username = "zhangzhi"
        api.log = mock.Mock()
        api._load_user_school_accounts = mock.Mock(return_value={})
        api.multi_get_all_accounts_status = mock.Mock(
            return_value={"success": True, "accounts": []}
        )
        add_calls = []

        def fake_add(
            username, password, tag=None, params=None, *, only_if_missing=False
        ):
            add_calls.append(
                {"username": username, "password": password, "tag": tag}
            )
            if username in api.accounts:
                return {
                    "success": False,
                    "reason": "already_exists",
                    "message": f"账号 {username} 已存在，未重复添加。",
                    "username": username,
                }
            api.accounts[username] = _FakeAccount(username)
            return {"success": True, "accounts": []}

        api.multi_add_account = mock.Mock(side_effect=fake_add)
        api.add_calls = add_calls
        return api

    @staticmethod
    def _b64(data):
        if isinstance(data, str):
            data = data.encode("utf-8")
        return base64.b64encode(data).decode("utf-8")

    @staticmethod
    def _xlsx(rows, title=None, merge_title=False):
        wb = openpyxl.Workbook()
        sh = wb.active
        sh.title = "模板"
        start_row = 1
        if title is not None:
            sh.cell(row=1, column=1, value=title)
            if merge_title:
                width = max(len(rows[0]) if rows else 4, 4)
                sh.merge_cells(
                    start_row=1, start_column=1, end_row=1, end_column=width
                )
            start_row = 2
        for row_offset, row in enumerate(rows):
            for col_offset, value in enumerate(row, start=1):
                sh.cell(
                    row=start_row + row_offset, column=col_offset, value=value
                )
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def _xlsx_b64(self, rows, title=None, merge_title=False):
        return self._b64(self._xlsx(rows, title=title, merge_title=merge_title))


class ImportFormatCompatTests(_ImportTestBase):
    def test_copy_template_xlsx_maps_by_header_and_ignores_serial(self):
        """副本模板：序号列必须忽略，账号/密码按表头正确定位。"""
        api = self._make_api()
        rows = [
            ["序号", "账号", "密码", "标记"],
            [1, VALID_A, "166354", "A组"],
            [2, VALID_B, "123456", ""],
        ]

        result = api.multi_import_accounts(
            "副本账号导入模板.xlsx",
            self._xlsx_b64(rows, title="莲峰账号导入表", merge_title=True),
        )

        self.assertTrue(result["success"])
        self.assertEqual(result["imported_count"], 2)
        self.assertEqual(
            api.add_calls,
            [
                {"username": VALID_A, "password": "166354", "tag": "A组"},
                {"username": VALID_B, "password": "123456", "tag": None},
            ],
        )
        self.assertNotIn("序号", api.accounts)
        self.assertNotIn("1", api.accounts)

    def test_title_row_containing_account_word_is_not_treated_as_header(self):
        """标题行自身含“账号/密码”二字时，也不能顶掉真正的表头行。"""
        api = self._make_api()
        rows = [
            ["账号密码统计表"],
            ["学号", "密码", "标记"],
            [VALID_A, "166354", "A组"],
        ]

        result = api.multi_import_accounts(
            "t.xlsx", self._xlsx_b64(rows, merge_title=False)
        )

        self.assertEqual(result["imported_count"], 1)
        self.assertEqual(api.add_calls[0]["username"], VALID_A)
        self.assertEqual(result["invalid_username_count"], 0)

    def test_copy_template_csv_same_mapping(self):
        api = self._make_api()
        text = f"序号,账号,密码,标记\n1,{VALID_A},166354,A组\n"

        result = api.multi_import_accounts("t.csv", self._b64(text))

        self.assertEqual(result["imported_count"], 1)
        self.assertEqual(
            api.add_calls[0],
            {"username": VALID_A, "password": "166354", "tag": "A组"},
        )

    def test_legacy_header_only_file(self):
        api = self._make_api()
        text = f"账号,密码,标记\n{VALID_A},166354,A组\n"

        result = api.multi_import_accounts("t.csv", self._b64(text))

        self.assertEqual(result["imported_count"], 1)
        self.assertEqual(api.add_calls[0]["tag"], "A组")

    def test_headerless_positional_file(self):
        api = self._make_api()
        text = f"{VALID_A},166354,A组\n{VALID_B},123456,\n"

        result = api.multi_import_accounts("t.csv", self._b64(text))

        self.assertEqual(result["imported_count"], 2)
        self.assertEqual(
            [call["tag"] for call in api.add_calls], ["A组", None]
        )

    def test_headerless_with_serial_column_infers_account_column(self):
        """无表头但第 1 列是序号：按数据内容推断第 2 列为学号列。"""
        api = self._make_api()
        text = f"1,{VALID_A},166354,A组\n2,{VALID_B},123456,B组\n"

        result = api.multi_import_accounts("t.csv", self._b64(text))

        self.assertEqual(result["imported_count"], 2)
        self.assertEqual(
            [call["username"] for call in api.add_calls], [VALID_A, VALID_B]
        )
        self.assertEqual(
            [call["password"] for call in api.add_calls], ["166354", "123456"]
        )
        self.assertNotIn("1", api.accounts)

    def test_headerless_keeps_leftmost_numeric_column_as_account(self):
        """无表头且账号列本身就是数字时，取最左的数字列当学号列。"""
        api = self._make_api()
        text = f"{VALID_A},123456\n{VALID_B},654321\n"

        result = api.multi_import_accounts("t.csv", self._b64(text))

        self.assertEqual(result["imported_count"], 2)
        self.assertEqual(
            [call["password"] for call in api.add_calls], ["123456", "654321"]
        )

    def test_roster_like_file_is_fully_rejected(self):
        """把“序号/姓名”名单当账号文件导入：整体拦下并完整上报。"""
        api = self._make_api()
        text = "序号,姓名\n1,朱燕姿\n2,石泳芝\n"

        result = api.multi_import_accounts("名单.csv", self._b64(text))

        self.assertTrue(result["success"])
        self.assertEqual(result["imported_count"], 0)
        self.assertEqual(result["total_rows"], 3)
        self.assertEqual(result["invalid_username_count"], 3)
        self.assertEqual(result["invalid_username"], ["序号", "1", "2"])
        self.assertEqual(result["invalid_password_count"], 0)
        self.assertEqual(api.accounts, {})
        self.assertIn("学号格式不对", result["message"])

    def test_chinese_password_is_rejected(self):
        api = self._make_api()
        text = f"账号,密码,标记\n{VALID_A},朱燕姿,A组\n"

        result = api.multi_import_accounts("t.csv", self._b64(text))

        self.assertEqual(result["imported_count"], 0)
        self.assertEqual(result["invalid_password_count"], 1)
        self.assertEqual(result["invalid_password"], [VALID_A])
        self.assertIn("密码包含中文", result["message"])
        self.assertEqual(api.accounts, {})

    def test_saved_chinese_password_is_rejected_on_fallback(self):
        """行内没密码时回退到已保存密码，回退值含中文同样要拦下。"""
        api = self._make_api()
        api._load_user_school_accounts.return_value = {
            VALID_A: {"password": "朱燕姿"}
        }
        text = f"账号,密码,标记\n{VALID_A},,\n"

        result = api.multi_import_accounts("t.csv", self._b64(text))

        self.assertEqual(result["imported_count"], 0)
        self.assertEqual(result["invalid_password_count"], 1)

    def test_missing_password_falls_back_to_saved_password(self):
        api = self._make_api()
        api._load_user_school_accounts.return_value = {
            VALID_A: {"password": "166354"}
        }
        text = f"账号,密码,标记\n{VALID_A},,\n"

        result = api.multi_import_accounts("t.csv", self._b64(text))

        self.assertEqual(result["imported_count"], 1)
        self.assertEqual(result["skipped_count"], 0)
        self.assertEqual(api.add_calls[0]["password"], "166354")

    def test_mixed_rows_report_every_bucket(self):
        api = self._make_api(existing=("2024000000001",))
        running_user = "2024000000002"
        base_add = api.multi_add_account.side_effect

        def fake_add(
            username, password, tag=None, params=None, *, only_if_missing=False
        ):
            if username == running_user:
                return {
                    "success": False,
                    "reason": "account_running",
                    "message": "正在执行任务",
                }
            return base_add(
                username,
                password,
                tag=tag,
                params=params,
                only_if_missing=only_if_missing,
            )

        api.multi_add_account = mock.Mock(side_effect=fake_add)
        text = (
            "账号,密码,标记\n"
            "2024000000003,pw,\n"       # 新增
            "2024000000001,pw,\n"       # 已存在
            "2024000000004,\n"          # 缺密码
            "2024000000003,pw,\n"       # 文件内重复
            "1,朱燕姿,\n"                # 非法学号
            "2024000000005,朱燕姿,\n"    # 密码含中文
            "2024000000002,pw,\n"       # 执行中
        )

        result = api.multi_import_accounts("mixed.csv", self._b64(text))

        self.assertEqual(result["total_rows"], 7)
        self.assertEqual(result["imported_count"], 1)
        self.assertEqual(result["already_exists_count"], 1)
        self.assertEqual(result["skipped_count"], 1)
        self.assertEqual(result["duplicate_in_file_count"], 1)
        self.assertEqual(result["invalid_username_count"], 1)
        self.assertEqual(result["invalid_password_count"], 1)
        self.assertEqual(result["skipped_running_count"], 1)
        self.assertEqual(result["failed_count"], 0)
        reasons = {item["username"]: item["reason"] for item in result["failed"]}
        self.assertEqual(reasons["1"], "invalid_username")
        self.assertEqual(reasons["2024000000005"], "invalid_password")

    @unittest.skipUnless(_HAS_XLWT, "xlwt 未安装，跳过 .xls 生成测试")
    def test_xls_with_serial_column_maps_by_header(self):
        api = self._make_api()
        wb = xlwt.Workbook()
        sh = wb.add_sheet("模板")
        for col, value in enumerate(["序号", "账号", "密码", "标记"]):
            sh.write(0, col, value)
        sh.write(1, 0, 1)
        sh.write(1, 1, VALID_A)
        sh.write(1, 2, "166354")
        sh.write(1, 3, "A组")
        output = io.BytesIO()
        wb.save(output)

        result = api.multi_import_accounts("t.xls", self._b64(output.getvalue()))

        self.assertEqual(result["imported_count"], 1)
        self.assertEqual(api.add_calls[0]["username"], VALID_A)
        self.assertEqual(api.add_calls[0]["password"], "166354")

    def test_numeric_cells_are_normalized_without_decimal_suffix(self):
        self.assertEqual(_normalize_import_cell(2024050101059.0), VALID_A)
        self.assertEqual(_normalize_import_cell(166354.0), "166354")
        self.assertEqual(_normalize_import_cell(" 1234 "), "1234")


class ImportTemplateTests(_ImportTestBase):
    def test_downloaded_template_matches_copy_layout(self):
        api = self._make_api()

        template = api.multi_download_import_template()

        self.assertTrue(template["success"])
        wb = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(template["content"]))
        )
        sh = wb.active
        self.assertEqual(sh.cell(row=1, column=1).value, "莲峰账号导入表")
        self.assertEqual(
            [sh.cell(row=2, column=col).value for col in range(1, 5)],
            ["序号", "账号", "密码", "标记"],
        )
        self.assertEqual(sh.cell(row=3, column=1).value, 1)

    def test_download_returns_repo_template_file_bytes_exactly(self):
        """下载模板必须直接返回仓库真实模板的字节，而不是现场重画一份。"""
        api = self._make_api()

        template = api.multi_download_import_template()

        expected = (FIXTURES_DIR / CANONICAL_TEMPLATE_FILENAME).read_bytes()
        self.assertTrue(template["success"])
        self.assertEqual(base64.b64decode(template["content"]), expected)

    def test_download_falls_back_to_generated_layout_when_fixture_missing(self):
        api = self._make_api()

        with mock.patch.object(
            api_runtime,
            "_IMPORT_TEMPLATE_RELATIVE_PATH",
            "tests/fixtures/__missing_template__.xlsx",
        ):
            template = api.multi_download_import_template()

        self.assertTrue(template["success"])
        sh = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(template["content"]))
        ).active
        self.assertEqual(sh.cell(row=1, column=1).value, "莲峰账号导入表")
        self.assertEqual(
            [sh.cell(row=2, column=col).value for col in range(1, 5)],
            ["序号", "账号", "密码", "标记"],
        )

    def test_downloaded_template_round_trips_with_filled_rows(self):
        api = self._make_api()
        template = api.multi_download_import_template()
        wb = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(template["content"]))
        )
        sh = wb.active
        sh.cell(row=3, column=2, value=VALID_A)
        sh.cell(row=3, column=3, value="166354")
        sh.cell(row=3, column=4, value="A组")
        output = io.BytesIO()
        wb.save(output)

        result = api.multi_import_accounts(
            "filled.xlsx", self._b64(output.getvalue())
        )

        self.assertEqual(result["imported_count"], 1)
        self.assertEqual(
            api.add_calls[-1],
            {"username": VALID_A, "password": "166354", "tag": "A组"},
        )

    def test_empty_downloaded_template_imports_nothing(self):
        api = self._make_api()
        template = api.multi_download_import_template()

        result = api.multi_import_accounts("empty.xlsx", template["content"])

        self.assertTrue(result["success"])
        self.assertEqual(result["total_rows"], 0)
        self.assertEqual(result["imported_count"], 0)

    def test_repo_keeps_real_template_files_with_original_names(self):
        """真实模板必须按原文件名保留在仓库，测试直接读它们而不是只看合成文件。"""
        for filename in (LEGACY_TEMPLATE_FILENAME, CANONICAL_TEMPLATE_FILENAME):
            with self.subTest(filename=filename):
                self.assertTrue(
                    (FIXTURES_DIR / filename).is_file(),
                    f"仓库缺少真实模板: {filename}",
                )

    def test_real_canonical_template_matches_generated_template_layout(self):
        """下载模板与仓库真实模板（school_accounts_template.xlsx）保持同一版式。"""
        api = self._make_api()
        generated = api.multi_download_import_template()
        generated_sheet = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(generated["content"]))
        ).active
        fixture_sheet = openpyxl.load_workbook(
            FIXTURES_DIR / CANONICAL_TEMPLATE_FILENAME
        ).active

        self.assertEqual(
            generated_sheet.cell(row=1, column=1).value,
            fixture_sheet.cell(row=1, column=1).value,
        )
        for column in range(1, 5):
            self.assertEqual(
                generated_sheet.cell(row=2, column=column).value,
                fixture_sheet.cell(row=2, column=column).value,
            )
        self.assertEqual(
            generated_sheet.cell(row=3, column=1).value,
            fixture_sheet.cell(row=3, column=1).value,
        )

    def test_real_canonical_template_empty_imports_nothing(self):
        api = self._make_api()
        content = (FIXTURES_DIR / CANONICAL_TEMPLATE_FILENAME).read_bytes()

        result = api.multi_import_accounts(
            CANONICAL_TEMPLATE_FILENAME, self._b64(content)
        )

        self.assertTrue(result["success"])
        self.assertEqual(result["total_rows"], 0)
        self.assertEqual(result["imported_count"], 0)
        self.assertEqual(result["invalid_username_count"], 0)

    def test_real_canonical_template_filled_rows_import(self):
        """真实副本模板填入账号后，序号列必须被忽略。"""
        api = self._make_api()
        wb = openpyxl.load_workbook(
            FIXTURES_DIR / CANONICAL_TEMPLATE_FILENAME
        )
        sh = wb.active
        sh.cell(row=3, column=2, value=VALID_A)
        sh.cell(row=3, column=3, value="166354")
        sh.cell(row=3, column=4, value="A组")
        sh.cell(row=4, column=2, value=VALID_B)
        sh.cell(row=4, column=3, value="123456")
        output = io.BytesIO()
        wb.save(output)

        result = api.multi_import_accounts(
            CANONICAL_TEMPLATE_FILENAME, self._b64(output.getvalue())
        )

        self.assertEqual(result["imported_count"], 2)
        self.assertEqual(result["invalid_username_count"], 0)
        self.assertEqual(
            [call["username"] for call in api.add_calls], [VALID_A, VALID_B]
        )
        self.assertEqual(api.add_calls[0]["tag"], "A组")
        self.assertNotIn("1", api.accounts)
        self.assertNotIn("序号", api.accounts)

    def test_real_legacy_template_appended_row_imports(self):
        """旧版“账号/密码/标记”模板仍要兼容（表头行被跳过）。"""
        api = self._make_api()
        wb = openpyxl.load_workbook(FIXTURES_DIR / LEGACY_TEMPLATE_FILENAME)
        sh = wb.active
        sh.append([VALID_A, "166354", "A组"])
        output = io.BytesIO()
        wb.save(output)

        result = api.multi_import_accounts(
            LEGACY_TEMPLATE_FILENAME, self._b64(output.getvalue())
        )

        self.assertEqual(result["imported_count"], 1)
        self.assertEqual(api.add_calls[0]["username"], VALID_A)
        self.assertEqual(api.add_calls[0]["password"], "166354")


class SchoolAccountValidatorTests(unittest.TestCase):
    def test_username_requires_four_plus_digits(self):
        for value in ("1234", "2024050101059", " 1234 ", 1234):
            with self.subTest(value=value):
                self.assertTrue(main_module.is_valid_school_username(value))
        for value in (
            "",
            None,
            "123",
            "abcd",
            "1234a",
            "序号",
            "１２３４",
            "2024-0501",
        ):
            with self.subTest(value=value):
                self.assertFalse(main_module.is_valid_school_username(value))

    def test_password_rejects_cjk_only(self):
        for value in ("166354", "abc123", "Abc#123", "", None):
            with self.subTest(value=value):
                self.assertTrue(main_module.is_valid_school_password(value))
        for value in ("朱燕姿", "abc朱", "密码123"):
            with self.subTest(value=value):
                self.assertFalse(main_module.is_valid_school_password(value))

    def test_multi_add_account_rejects_before_touching_session(self):
        api = object.__new__(main_module.Api)

        bad_username = main_module.Api.multi_add_account(api, "1", "pw", None)
        self.assertFalse(bad_username["success"])
        self.assertEqual(bad_username["reason"], "invalid_username")

        bad_password = main_module.Api.multi_add_account(
            api, VALID_A, "朱燕姿", None
        )
        self.assertFalse(bad_password["success"])
        self.assertEqual(bad_password["reason"], "invalid_password")


class SchoolAccountStorageValidationTests(unittest.TestCase):
    def setUp(self):
        self._temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self._temp_dir.cleanup)
        self.root = Path(self._temp_dir.name) / "school_accounts"
        (self.root / "user_accounts").mkdir(parents=True)

    def _api(self):
        api = object.__new__(main_module.Api)
        api.user_dir = str(self.root)
        api.auth_username = "alice"
        api.log = mock.Mock()
        return api

    def test_load_filters_invalid_usernames(self):
        (self.root / "user_accounts" / "alice.json").write_text(
            json.dumps(
                {
                    "1": {"password": "朱燕姿", "ua": ""},
                    "序号": {"password": "姓名", "ua": ""},
                    VALID_A: {"password": "166354", "ua": "ua"},
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        api = self._api()

        with mock.patch.multiple(
            main_module,
            SCHOOL_ACCOUNTS_DIR=str(self.root),
            auth_system=mock.Mock(),
            create=True,
        ):
            main_module.auth_system.check_permission.return_value = False
            accounts = api._load_user_school_accounts("alice")

        self.assertEqual(list(accounts.keys()), [VALID_A])

    def test_save_drops_invalid_usernames_and_ini_files(self):
        api = self._api()

        with mock.patch.multiple(
            main_module,
            SCHOOL_ACCOUNTS_DIR=str(self.root),
            auth_system=mock.Mock(),
            create=True,
        ):
            api._save_user_school_accounts(
                "alice",
                {
                    "1": {"password": "朱燕姿", "ua": ""},
                    "序号": {"password": "姓名", "ua": ""},
                    VALID_A: {
                        "password": "166354",
                        "ua": "ua",
                        "overdue_count": 0,
                        "completed_count": 0,
                    },
                },
            )

        data = json.loads(
            (self.root / "user_accounts" / "alice.json").read_text(
                encoding="utf-8"
            )
        )
        self.assertEqual(list(data.keys()), [VALID_A])
        self.assertFalse((self.root / "1.ini").exists())
        self.assertFalse((self.root / "序号.ini").exists())
        self.assertTrue((self.root / f"{VALID_A}.ini").exists())

    def test_admin_aggregation_skips_invalid_legacy_ini(self):
        (self.root / "1.ini").write_text(
            "[Config]\npassword = 朱燕姿\n", encoding="utf-8"
        )
        (self.root / "序号.ini").write_text(
            "[Config]\npassword = 姓名\n", encoding="utf-8"
        )
        (self.root / f"{VALID_A}.ini").write_text(
            "[Config]\npassword = 166354\n", encoding="utf-8"
        )
        api = self._api()

        with mock.patch.multiple(
            main_module,
            SCHOOL_ACCOUNTS_DIR=str(self.root),
            auth_system=mock.Mock(),
            create=True,
        ):
            main_module.auth_system.check_permission.return_value = True
            accounts = api._load_user_school_accounts("admin")

        self.assertIn(VALID_A, accounts)
        self.assertNotIn("1", accounts)
        self.assertNotIn("序号", accounts)

    def test_save_config_refuses_invalid_username(self):
        api = self._api()

        with mock.patch.multiple(
            main_module,
            SCHOOL_ACCOUNTS_DIR=str(self.root),
            auth_system=mock.Mock(),
            create=True,
        ):
            saved = api._save_config("1", "pw")

        self.assertFalse(saved)
        self.assertFalse((self.root / "1.ini").exists())


class ImportWiringTests(unittest.TestCase):
    """前端与管理端路由必须把新的分类/校验展示、接线到位。"""

    def _project_root(self):
        return Path(main_module.__file__).resolve().parent

    def test_admin_routes_validate_username_and_password(self):
        source = (self._project_root() / "web_server.py").read_text(
            encoding="utf-8"
        )
        self.assertGreaterEqual(
            source.count("is_valid_school_username(school_username)"), 2
        )
        self.assertGreaterEqual(
            source.count("is_valid_school_password(password)"), 2
        )
        self.assertIn("学号必须是至少4位数字", source)
        self.assertIn("密码不能包含中文", source)

    def test_legacy_import_modal_renders_invalid_sections(self):
        source = (self._project_root() / "scripts" / "main.js").read_text(
            encoding="utf-8"
        )
        self.assertIn("result.invalid_username", source)
        self.assertIn("result.invalid_password", source)
        self.assertIn("学号格式不对", source)
        self.assertIn("密码包含中文", source)


if __name__ == "__main__":
    unittest.main()
